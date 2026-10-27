"""End-to-end tests for the signing ceremony API."""
from django.test import override_settings

import base64
import http.server
import io
import threading
from datetime import timedelta

import pymupdf
import pytest
from PIL import Image
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient

from .hashing import GENESIS_HASH, verify_chain
from .trust import assess_photo, normalize_photo
from .models import AuditEvent, Document, SigningSession

# 1x1 transparent PNG — stands in for a drawn signature / camera photo.
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)
SIG_DATA_URL = "data:image/png;base64," + base64.b64encode(TINY_PNG).decode()


def make_pdf_bytes(text="Test document") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


def upload_document(client, title="Contract", email="sender@example.com"):
    pdf = SimpleUploadedFile(
        "contract.pdf", make_pdf_bytes(), content_type="application/pdf"
    )
    return client.post(
        "/api/documents/",
        {"title": title, "sender_email": email, "file": pdf},
        format="multipart",
    )


@pytest.mark.django_db
def test_upload_document_creates_pending_document():
    client = APIClient()
    response = upload_document(client)
    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "pending"
    assert body["title"] == "Contract"
    assert Document.objects.count() == 1


@pytest.mark.django_db
def test_upload_rejects_non_pdf():
    client = APIClient()
    fake = SimpleUploadedFile("evil.txt", b"not a pdf", content_type="text/plain")
    response = client.post(
        "/api/documents/",
        {"title": "X", "sender_email": "s@example.com", "file": fake},
        format="multipart",
    )
    assert response.status_code == 400
    assert "PDF" in response.json()["error"]


@pytest.mark.django_db
def test_upload_requires_title_and_email():
    client = APIClient()
    pdf = SimpleUploadedFile("c.pdf", make_pdf_bytes(), content_type="application/pdf")
    response = client.post("/api/documents/", {"file": pdf}, format="multipart")
    assert response.status_code == 400


@pytest.mark.django_db
def test_create_session_returns_token_and_signing_url():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    response = client.post(
        f"/api/documents/{doc_id}/sessions/",
        {"signer_name": "Jane", "signer_email": "jane@example.com"},
        format="json",
    )
    assert response.status_code == 201
    body = response.json()
    assert body["token"]
    assert body["signing_url"].endswith(f"/s/{body['token']}")
    assert body["usable"] is True
    # session.created is the genesis event of the chain
    event = AuditEvent.objects.get(session_id=body["id"])
    assert event.event_type == "session.created"
    assert event.prev_hash == GENESIS_HASH


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_create_session_requires_signer_email():
    from django.core import mail

    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    assert (
        client.post(
            f"/api/documents/{doc_id}/sessions/",
            {"signer_name": "Jane"},
            format="json",
        ).status_code
        == 400
    )
    assert (
        client.post(
            f"/api/documents/{doc_id}/sessions/",
            {"signer_name": "Jane", "signer_email": "  "},
            format="json",
        ).status_code
        == 400
    )
    assert SigningSession.objects.count() == 0


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_signer_invite_attaches_unsigned_pdf():
    # Email-first delivery: the invite alone must be everything the
    # signer needs — signing link plus the document to sign.
    from django.core import mail

    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    assert (
        client.post(
            f"/api/documents/{doc_id}/sessions/",
            {"signer_name": "Jane", "signer_email": "jane@example.com"},
            format="json",
        ).status_code
        == 201
    )
    invites = [m for m in mail.outbox if m.subject.startswith("Please sign:")]
    assert len(invites) == 1
    invite = invites[0]
    assert invite.to == ["jane@example.com"]
    assert "/s/" in invite.body
    assert len(invite.attachments) == 1
    filename, _content, mime = invite.attachments[0]
    assert filename == "contract.pdf"
    assert mime == "application/pdf"


@pytest.mark.django_db
def test_full_ceremony_signs_seals_and_verifies():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]

    response = client.post(
        f"/api/sessions/{token}/sign/",
        {
            "signature": SIG_DATA_URL,
            "photo": SIG_DATA_URL,
            "photo_consent": True,
            "device_info": {"ua": "test", "platform": "test"},
        },
        format="json",
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "signed"
    assert body["chain_valid"] is True
    assert body["document"]["status"] == "signed"
    assert body["document"]["signed_file"]

    doc = Document.objects.get(pk=doc_id)
    assert doc.signed_file
    # Sealed PDF is a real PDF, larger than the original (certificate page).
    assert doc.signed_file.size > len(make_pdf_bytes())

    session = SigningSession.objects.get(token=token)
    assert session.status == "signed"
    types = [e.event_type for e in session.events.order_by("created_at")]
    assert types == ["session.created", "signature.submitted", "document.sealed"]
    assert verify_chain(session) is True


@pytest.mark.django_db
def test_double_sign_is_rejected():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    payload = {"signature": SIG_DATA_URL}
    assert client.post(f"/api/sessions/{token}/sign/", payload, format="json").status_code == 200
    second = client.post(f"/api/sessions/{token}/sign/", payload, format="json")
    assert second.status_code == 410


@pytest.mark.django_db
def test_expired_session_cannot_sign():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    session = SigningSession.objects.get(token=token)
    session.expires_at = timezone.now() - timedelta(seconds=1)
    session.save()

    response = client.post(
        f"/api/sessions/{token}/sign/", {"signature": SIG_DATA_URL}, format="json"
    )
    assert response.status_code == 410


@pytest.mark.django_db
def test_tampered_chain_is_detected():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    client.post(
        f"/api/sessions/{token}/sign/", {"signature": SIG_DATA_URL}, format="json"
    )
    session = SigningSession.objects.get(token=token)
    assert verify_chain(session) is True

    # Attacker rewrites a payload in place.
    event = session.events.order_by("created_at").first()
    event.payload = {"tampered": True}
    event.save()

    assert verify_chain(session) is False
    # The public audit endpoint reports it too.
    events_resp = client.get(f"/api/sessions/{token}/events/")
    assert events_resp.json()["chain_valid"] is False


@pytest.mark.django_db
def test_sign_requires_a_signature_image():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    response = client.post(f"/api/sessions/{token}/sign/", {}, format="json")
    assert response.status_code == 400


@pytest.mark.django_db
def test_sign_captures_signer_email():
    # Email-first delivery requires an address at creation; the sign
    # payload can still correct it (e.g. the invite went to a stale
    # address) and the corrected address is what the session keeps.
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "old@example.com"}, format="json"
    ).json()["token"]
    response = client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "signer_email": "jane@example.com"},
        format="json",
    )
    assert response.status_code == 200
    session = SigningSession.objects.get(token=token)
    assert session.signer_email == "jane@example.com"


@pytest.mark.django_db
def test_session_document_serves_pdf_to_token_holder():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    response = client.get(f"/api/sessions/{token}/document/")
    assert response.status_code == 200
    assert response["Content-Type"] == "application/pdf"


def make_docx_bytes(text="Test Agreement") -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("Test Agreement", 1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
def test_docx_upload_converts_to_pdf():
    client = APIClient()
    upload = SimpleUploadedFile(
        "agreement.docx",
        make_docx_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response = client.post(
        "/api/documents/",
        {"title": "Agreement", "sender_email": "s@example.com", "file": upload},
        format="multipart",
    )
    assert response.status_code == 201
    body = response.json()
    assert body["source"] == "upload"
    assert body["original_filename"] == "agreement.docx"

    doc = Document.objects.get(pk=body["id"])
    with open(doc.original_file.path, "rb") as f:
        assert f.read(5) == b"%PDF-"


@pytest.mark.django_db
def test_upload_rejects_unsupported_extension():
    client = APIClient()
    upload = SimpleUploadedFile("notes.txt", b"hello", content_type="text/plain")
    response = client.post(
        "/api/documents/",
        {"title": "X", "sender_email": "s@example.com", "file": upload},
        format="multipart",
    )
    assert response.status_code == 400
    assert "PDF, DOC, or DOCX" in response.json()["error"]


@pytest.mark.django_db
@override_settings(URL_FETCH_ALLOW_PRIVATE=True)
def test_from_url_renders_page_to_pdf():
    """Real end-to-end: a local HTTP server stands in for the web page.

    The SSRF guard (issue #39) refuses private addresses; this test is
    the local-development case the URL_FETCH_ALLOW_PRIVATE hatch is for.
    """
    html = b"<html><body><h1>Terms of Service</h1><p>These are the terms.</p></body></html>"

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write(html)

        def log_message(self, *args):
            pass

    server = http.server.HTTPServer(("127.0.0.1", 0), Handler)
    port = server.server_address[1]
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        client = APIClient()
        response = client.post(
            "/api/documents/from-url/",
            {
                "title": "ToS",
                "sender_email": "s@example.com",
                "url": f"http://127.0.0.1:{port}/terms",
            },
            format="json",
        )
    finally:
        server.shutdown()

    assert response.status_code == 201
    body = response.json()
    assert body["source"] == "url"
    assert body["source_url"].endswith("/terms")

    doc = Document.objects.get(pk=body["id"])
    with open(doc.original_file.path, "rb") as f:
        assert f.read(5) == b"%PDF-"


@pytest.mark.django_db
def test_from_url_rejects_bad_scheme():
    client = APIClient()
    response = client.post(
        "/api/documents/from-url/",
        {"title": "X", "sender_email": "s@example.com", "url": "ftp://example.com/x"},
        format="json",
    )
    assert response.status_code == 400


# ---------------------------------------------------------------------------
# Phase 3: trust hardening
# ---------------------------------------------------------------------------


def _exif_jpeg(dt_str=None, offset=None, orientation=None, size=(100, 50)):
    """Build a JPEG with controlled EXIF tags for anti-spoof tests."""
    img = Image.new("RGB", size, (180, 40, 40))
    exif = Image.Exif()
    if dt_str:
        exif[36867] = dt_str  # DateTimeOriginal
    if offset:
        exif[36881] = offset  # OffsetTimeOriginal
    if orientation:
        exif[274] = orientation
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    return buf.getvalue()


def _exif_str(dt):
    return dt.strftime("%Y:%m:%d %H:%M:%S")


def test_exif_match_gives_full_assurance():
    now = timezone.now()
    photo = _exif_jpeg(_exif_str(now), offset="+00:00")
    report = assess_photo(photo, now, tz_offset_minutes=0)
    assert report["assurance"] == "full"
    assert report["reason"] == "timestamp_match"
    assert report["exif_present"] is True


def test_exif_match_via_client_timezone_offset():
    # Naive EXIF (no offset tag); client reports UTC-5 (Chicago, offset 300).
    now = timezone.now()
    local = now - timedelta(hours=5)
    photo = _exif_jpeg(_exif_str(local))
    report = assess_photo(photo, now, tz_offset_minutes=300)
    assert report["assurance"] == "full"
    assert report["reason"] == "timestamp_match"


def test_exif_missing_gives_reduced_assurance():
    photo = _exif_jpeg()
    report = assess_photo(photo, timezone.now(), tz_offset_minutes=0)
    assert report["assurance"] == "reduced"
    assert report["reason"] == "exif_timestamp_missing"


def test_exif_mismatch_gives_reduced_assurance():
    now = timezone.now()
    old = now - timedelta(hours=2)
    photo = _exif_jpeg(_exif_str(old), offset="+00:00")
    report = assess_photo(photo, now, tz_offset_minutes=0)
    assert report["assurance"] == "reduced"
    assert report["reason"] == "exif_timestamp_mismatch"
    assert report["delta_seconds"] == pytest.approx(7200, abs=5)


def test_exif_unparseable_gives_reduced_assurance():
    photo = _exif_jpeg("not-a-date", offset="+00:00")
    report = assess_photo(photo, timezone.now(), tz_offset_minutes=0)
    assert report["assurance"] == "reduced"
    assert report["reason"] == "exif_timestamp_unparseable"


def test_exif_naive_without_zone_gives_reduced_assurance():
    now = timezone.now()
    photo = _exif_jpeg(_exif_str(now))
    report = assess_photo(photo, now, tz_offset_minutes=None)
    assert report["assurance"] == "reduced"
    assert report["reason"] == "exif_timezone_unknown"


def test_exif_assessment_never_raises_on_garbage():
    report = assess_photo(b"definitely not an image", timezone.now(), 0)
    assert report["assurance"] == "reduced"


def test_normalize_photo_applies_exif_orientation():
    # Orientation 6 = stored rotated 90° CW; output must be transposed.
    photo = _exif_jpeg(orientation=6, size=(100, 50))
    out = normalize_photo(photo)
    assert Image.open(io.BytesIO(out)).size == (50, 100)


def _sign_with_photo(client, photo_bytes, consent=True, device_info=None):
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    payload = {
        "signature": SIG_DATA_URL,
        "photo": "data:image/jpeg;base64,"
        + base64.b64encode(photo_bytes).decode(),
        "device_info": device_info or {},
    }
    if consent:
        payload["photo_consent"] = True
    response = client.post(f"/api/sessions/{token}/sign/", payload, format="json")
    return doc_id, token, response


@pytest.mark.django_db
def test_photo_requires_consent():
    client = APIClient()
    _, _, response = _sign_with_photo(client, _exif_jpeg(), consent=False)
    assert response.status_code == 400
    assert "photo_consent" in response.json()["error"]


@pytest.mark.django_db
def test_submit_records_device_evidence_and_photo_assurance():
    client = APIClient()
    now = timezone.now()
    photo = _exif_jpeg(_exif_str(now), offset="+00:00")
    doc_id, token, response = _sign_with_photo(
        client,
        photo,
        device_info={
            "userAgent": "test-ua",
            "platform": "test-os",
            "language": "en",
            "screenWidth": 390,
            "screenHeight": 844,
            "timezoneOffsetMinutes": 0,
        },
    )
    assert response.status_code == 200

    session = SigningSession.objects.get(token=token)
    assert session.device_evidence["user_agent"] == "test-ua"
    assert session.device_evidence["screen"] == {"width": 390, "height": 844}
    assert session.device_evidence["timezone_offset_minutes"] == 0

    submitted = session.events.get(event_type="signature.submitted")
    assert submitted.payload["photo_consent"] is True
    assurance = submitted.payload["photo_assurance"]
    assert assurance["assurance"] == "full"
    assert assurance["reason"] == "timestamp_match"

    # The certificate page carries the assurance line.
    doc = Document.objects.get(pk=doc_id)
    pdf = pymupdf.open(stream=doc.signed_file.read())
    assert "Witness photo assurance: full" in pdf[-1].get_text()


@pytest.mark.django_db
def test_sealed_pdf_carries_free_tier_watermark():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    response = client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    )
    assert response.status_code == 200
    doc = Document.objects.get(pk=doc_id)
    pdf = pymupdf.open(stream=doc.signed_file.read())
    assert len(pdf) == 2  # original page + certificate page
    for page in pdf:
        assert "Signed with SignSeen" in page.get_text()


@pytest.mark.django_db
def test_cleanup_expired_marks_stale_sessions():
    from django.core.management import call_command

    client = APIClient()
    stale_doc_id = upload_document(client, email="old@example.com").json()["id"]
    stale_token = client.post(
        f"/api/documents/{stale_doc_id}/sessions/", {"signer_name": "Old", "signer_email": "old@example.com"}, format="json"
    ).json()["token"]
    fresh_doc_id = upload_document(client, title="Fresh", email="new@example.com").json()["id"]
    fresh_token = client.post(
        f"/api/documents/{fresh_doc_id}/sessions/", {"signer_name": "New", "signer_email": "new@example.com"}, format="json"
    ).json()["token"]
    stale = SigningSession.objects.get(token=stale_token)
    stale.expires_at = timezone.now() - timedelta(seconds=1)
    stale.save(update_fields=["expires_at"])

    call_command("cleanup_expired")

    stale.refresh_from_db()
    fresh = SigningSession.objects.get(token=fresh_token)
    assert stale.status == SigningSession.Status.EXPIRED
    assert fresh.status == SigningSession.Status.PENDING
    # The fresh document stays pending: it still has a live session.
    assert Document.objects.get(pk=fresh_doc_id).status == Document.Status.PENDING


# ---------------------------------------------------------------------------
# Phase 4: freemium (accounts, enforcement, watermark gating, email)
# ---------------------------------------------------------------------------

from django.test import override_settings

from apps.accounts.models import Account


def _make_account(email="owner@example.com", tier=Account.Tier.FREE):
    account = Account.objects.create_user(
        email=email, password="longpassword1"
    )
    if tier != Account.Tier.FREE:
        account.tier = tier
        account.save(update_fields=["tier"])
    return account


def _auth_client(account):
    client = APIClient()
    token = client.post(
        "/api/auth/login/",
        {"email": account.email, "password": "longpassword1"},
        format="json",
    ).json()["token"]
    client.credentials(HTTP_AUTHORIZATION=f"Token {token}")
    return client


@pytest.mark.django_db
def test_free_tier_allows_only_one_pending_document():
    client = APIClient()
    assert upload_document(client).status_code == 201
    second = upload_document(client, title="Second")
    assert second.status_code == 409
    assert "1 document pending" in second.json()["error"]


@pytest.mark.django_db
def test_signed_document_frees_the_pending_slot():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    assert upload_document(client, title="Next").status_code == 201


@pytest.mark.django_db
def test_account_documents_get_owner_and_share_the_limit():
    account = _make_account()
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    assert Document.objects.get(pk=doc_id).owner == account
    # The account's pending doc blocks a second one, even anonymously
    # addressed to the same email (identity = account when logged in;
    # anonymous path counts by email and sees the same pending doc).
    assert upload_document(client, email=account.email).status_code == 409


@pytest.mark.django_db
def test_premium_bypasses_pending_limit():
    account = _make_account(email="pro@example.com", tier=Account.Tier.PREMIUM)
    client = _auth_client(account)
    assert upload_document(client, email=account.email).status_code == 201
    assert upload_document(client, title="Two", email=account.email).status_code == 201


@pytest.mark.django_db
@override_settings(FREE_STORAGE_BYTES=1)
def test_free_storage_cap_is_enforced():
    client = APIClient()
    response = upload_document(client)
    assert response.status_code == 409
    assert "Storage limit" in response.json()["error"]


@pytest.mark.django_db
def test_second_signing_session_is_rejected():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    assert client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "One", "signer_email": "one@example.com"}, format="json"
    ).status_code == 201
    second = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Two", "signer_email": "two@example.com"}, format="json"
    )
    assert second.status_code == 409


@pytest.mark.django_db
def test_my_documents_requires_auth_and_reports_usage():
    account = _make_account()
    client = _auth_client(account)
    assert APIClient().get("/api/documents/mine/").status_code == 401
    doc_id = upload_document(client, email=account.email).json()["id"]
    body = client.get("/api/documents/mine/").json()
    assert body["tier"] == "free"
    assert [d["id"] for d in body["documents"]] == [doc_id]
    assert body["storage_used_bytes"] > 0
    assert body["storage_limit_bytes"] == 100 * 1024 * 1024


def _sign_and_get_pdf_text(client, doc_id):
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    doc = Document.objects.get(pk=doc_id)
    return pymupdf.open(stream=doc.signed_file.read())[-1].get_text(), doc


@pytest.mark.django_db
def test_free_document_is_watermarked_premium_is_not():
    # Free (anonymous): watermark present.
    anon = APIClient()
    free_doc_id = upload_document(anon).json()["id"]
    _, _ = _sign_and_get_pdf_text(anon, free_doc_id)

    # Premium (account): no watermark anywhere in the sealed PDF.
    account = _make_account(email="pro2@example.com", tier=Account.Tier.PREMIUM)
    client = _auth_client(account)
    pro_doc_id = upload_document(client, email=account.email).json()["id"]
    token = client.post(
        f"/api/documents/{pro_doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    doc = Document.objects.get(pk=pro_doc_id)
    pdf = pymupdf.open(stream=doc.signed_file.read())
    for page in pdf:
        assert "free tier" not in page.get_text()


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_sealed_email_goes_to_sender_and_signer():
    from django.core import mail

    client = APIClient()
    doc_id = upload_document(client, email="sender@example.com").json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {
            "signature": SIG_DATA_URL,
            "device_info": {},
            "signer_email": "signer@example.com",
        },
        format="json",
    ).status_code == 200
    # The invite (to the creation-time address) is also in the outbox;
    # the sealed copies are what this test owns.
    sealed = [m for m in mail.outbox if m.subject.startswith("Signed:")]
    recipients = sorted(m.to[0] for m in sealed)
    assert recipients == ["sender@example.com", "signer@example.com"]
    for message in sealed:
        assert len(message.attachments) == 1
        assert "Audit seal" in message.body


@pytest.mark.django_db
def test_pending_limit_is_keyed_on_email_across_account_and_anonymous():
    # Anonymous send holds the email's single pending slot...
    anon = APIClient()
    assert upload_document(anon, email="same@example.com").status_code == 201
    # ...so the same email, now logged in, is still blocked.
    account = _make_account(email="same@example.com")
    authed = _auth_client(account)
    assert upload_document(authed, email="same@example.com").status_code == 409
    # And the reverse: an account's pending doc blocks anonymous use of it.
    account2 = _make_account(email="other@example.com")
    authed2 = _auth_client(account2)
    assert upload_document(authed2, email="other@example.com").status_code == 201
    assert upload_document(anon, email="other@example.com").status_code == 409


@pytest.mark.django_db
def test_login_claims_previously_anonymous_documents():
    # Sent anonymously (ownerless, keyed by email only)...
    anon = APIClient()
    doc_id = upload_document(anon, email="claim@example.com").json()["id"]
    assert Document.objects.get(pk=doc_id).owner is None
    # ...then the sender creates an account with the same email: the
    # document is claimed and appears on their dashboard.
    client = APIClient()
    token = client.post(
        "/api/auth/signup/",
        {"email": "claim@example.com", "password": "longpassword1"},
        format="json",
    ).json()["token"]
    client.credentials(HTTP_AUTHORIZATION=f"Token {token}")
    assert Document.objects.get(pk=doc_id).owner is not None
    body = client.get("/api/documents/mine/").json()
    assert [d["id"] for d in body["documents"]] == [doc_id]


@pytest.mark.django_db
def test_signer_archive_shows_documents_they_signed():
    # Someone else sends; our signer signs and leaves their email.
    sender = APIClient()
    doc_id = upload_document(sender, email="sender2@example.com").json()["id"]
    token = sender.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Sig", "signer_email": "sig@example.com"}, format="json"
    ).json()["token"]
    assert sender.post(
        f"/api/sessions/{token}/sign/",
        {
            "signature": SIG_DATA_URL,
            "device_info": {},
            "signer_email": "sig@example.com",
        },
        format="json",
    ).status_code == 200
    # The signer creates an account with that email post-signing.
    client = APIClient()
    auth = client.post(
        "/api/auth/signup/",
        {"email": "sig@example.com", "password": "longpassword1"},
        format="json",
    ).json()["token"]
    client.credentials(HTTP_AUTHORIZATION=f"Token {auth}")
    body = client.get("/api/documents/mine/").json()
    assert [d["id"] for d in body["signed_documents"]] == [doc_id]
    assert body["documents"] == []


# ---------------------------------------------------------------------------
# Stale sender tokens must not break public / signer endpoints.
#
# A browser keeps its sender token in localStorage after the server has
# deleted it (logout on another device, password reset, DB reset). DRF
# TokenAuthentication rejects unknown tokens before permission checks,
# so these endpoints either ignore the Authorization header entirely
# (token-addressed signer endpoints, public document view) or degrade
# an invalid token to anonymous (optionally-authenticated sender ones).
# ---------------------------------------------------------------------------

STALE = "Token stale-token-no-longer-in-db"


def _stale_client():
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=STALE)
    return client


@pytest.mark.django_db
def test_signer_endpoints_ignore_stale_sender_token():
    anon = APIClient()
    doc_id = upload_document(anon).json()["id"]
    token = anon.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]

    stale = _stale_client()
    session = stale.get(f"/api/sessions/{token}/")
    assert session.status_code == 200
    assert session.json()["usable"] is True

    document = stale.get(f"/api/sessions/{token}/document/")
    assert document.status_code == 200

    events = stale.get(f"/api/sessions/{token}/events/")
    assert events.status_code == 200
    assert events.json()["chain_valid"] is True

    signed = stale.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    )
    assert signed.status_code == 200
    assert signed.json()["status"] == "signed"


@pytest.mark.django_db
def test_public_document_view_ignores_stale_sender_token():
    anon = APIClient()
    doc_id = upload_document(anon).json()["id"]
    response = _stale_client().get(f"/api/documents/{doc_id}/")
    assert response.status_code == 200
    assert response.json()["id"] == doc_id


@pytest.mark.django_db
def test_anonymous_upload_with_stale_token_degrades_to_anonymous():
    stale = _stale_client()
    response = upload_document(stale)
    assert response.status_code == 201
    assert Document.objects.get(pk=response.json()["id"]).owner is None
    # Session creation on that document works the same way.
    session = stale.post(
        f"/api/documents/{response.json()['id']}/sessions/",
        {"signer_name": "Jane", "signer_email": "jane@example.com"},
        format="json",
    )
    assert session.status_code == 201


# ---------------------------------------------------------------------------
# Completion email step (scope §3.3/§3.10): the signer leaves an email
# after sealing to receive a copy; that email is what links the signed
# document to a later account's "Documents you've signed" archive.
# ---------------------------------------------------------------------------


def _sign_document(client, email="sender@example.com"):
    doc_id = upload_document(client, email=email).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    return doc_id, token


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_signer_email_at_completion_sends_copy_and_links_archive():
    from django.core import mail

    client = APIClient()
    doc_id, token = _sign_document(client)
    # Seal-time mail went to the sender and the invited signer (the
    # session now carries an email from creation under email-first
    # delivery; the earlier invite is also in the outbox); the
    # completion endpoint can still redirect the copy to a different
    # address.
    assert [m.to for m in mail.outbox if m.subject.startswith("Signed:")] == [
        ["sender@example.com"],
        ["jane@example.com"],
    ]

    # A stale sender Authorization header must not matter here either.
    stale = APIClient()
    stale.credentials(HTTP_AUTHORIZATION="Token stale-token-no-longer-in-db")
    response = stale.post(
        f"/api/sessions/{token}/signer-email/",
        {"email": "jane-copy@example.com"},
        format="json",
    )
    assert response.status_code == 200
    session = SigningSession.objects.get(token=token)
    assert session.signer_email == "jane-copy@example.com"
    assert [m.to for m in mail.outbox][-1] == ["jane-copy@example.com"]
    assert session.events.filter(event_type="signer.email_provided").exists()
    assert verify_chain(session) is True

    # An account created with that email now archives the document.
    account_client = APIClient()
    auth = account_client.post(
        "/api/auth/signup/",
        {"email": "jane-copy@example.com", "password": "longpassword1"},
        format="json",
    ).json()["token"]
    account_client.credentials(HTTP_AUTHORIZATION=f"Token {auth}")
    body = account_client.get("/api/documents/mine/").json()
    assert [d["id"] for d in body["signed_documents"]] == [doc_id]


@pytest.mark.django_db
def test_signer_email_rejected_before_signing_and_when_invalid():
    client = APIClient()
    doc_id = upload_document(client).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    pending = client.post(
        f"/api/sessions/{token}/signer-email/",
        {"email": "jane@example.com"},
        format="json",
    )
    assert pending.status_code == 409

    _, signed_token = _sign_document(APIClient(), email="other@example.com")
    bad = client.post(
        f"/api/sessions/{signed_token}/signer-email/",
        {"email": "not-an-email"},
        format="json",
    )
    assert bad.status_code == 400
    # The rejected address leaves the invite-time address untouched.
    assert (
        SigningSession.objects.get(token=signed_token).signer_email
        == "jane@example.com"
    )


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Custom branding payload + outputs (issue #10 / scope §3.6)
# ---------------------------------------------------------------------------


def _branded_account(email="branded@example.com"):
    account = _make_account(email=email, tier=Account.Tier.PREMIUM)
    account.brand_name = "Acme Legal"
    account.brand_color = "#123abc"
    account.save(update_fields=["brand_name", "brand_color"])
    return account


@pytest.mark.django_db
def test_session_payload_exposes_owner_branding():
    account = _branded_account()
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    body = APIClient().get(f"/api/sessions/{token}/").json()
    assert body["branding"] == {"name": "Acme Legal", "color": "#123abc", "logo": None}


@pytest.mark.django_db
def test_session_payload_branding_empty_for_anonymous_and_unbranded():
    anon = APIClient()
    doc_id = upload_document(anon).json()["id"]
    token = anon.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    body = APIClient().get(f"/api/sessions/{token}/").json()
    assert body["branding"] == {"name": "", "color": "", "logo": None}

    # An owned document whose account set no branding gets the same empties.
    account = _make_account(email="plain@example.com", tier=Account.Tier.PREMIUM)
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    body = APIClient().get(f"/api/sessions/{token}/").json()
    assert body["branding"] == {"name": "", "color": "", "logo": None}


@pytest.mark.django_db
def test_sealed_pdf_certificate_shows_brand_name():
    account = _branded_account(email="pdf-brand@example.com")
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    text, _doc = _sign_and_get_pdf_text(client, doc_id)
    assert "via Acme Legal" in text


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_sealed_email_opens_with_brand_name():
    from django.core import mail

    account = _branded_account(email="mail-brand@example.com")
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    token = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "J", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    assert client.post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    sealed = [m for m in mail.outbox if m.subject.startswith("Signed:")]
    assert sealed
    for message in sealed:
        assert message.body.startswith("Acme Legal — signed document via SignSeen")

# ---------------------------------------------------------------------------
# Signing alerts (issue #9): viewed + expiring-soon emails. Premium
# owners with alerts enabled only; the record is the audit event +
# the session timestamp fields.
# ---------------------------------------------------------------------------


def _owned_session(account_client, email):
    doc_id = upload_document(account_client, email=email).json()["id"]
    token = account_client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()["token"]
    return doc_id, token


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_viewed_alert_fires_exactly_once_for_premium_owner():
    from django.core import mail

    account = _make_account(email="prem@example.com", tier=Account.Tier.PREMIUM)
    client = _auth_client(account)
    _, token = _owned_session(client, account.email)

    anon = APIClient()
    assert anon.get(f"/api/sessions/{token}/").status_code == 200
    session = SigningSession.objects.get(token=token)
    assert session.viewed_at is not None
    assert session.events.filter(event_type="signing.viewed").count() == 1
    assert verify_chain(session) is True
    # The invite ("Please sign:") also sits in the outbox; filter to
    # alerts, which is what this test owns.
    alerts = [m for m in mail.outbox if m.subject.startswith("Viewed:")]
    assert [m.to for m in alerts] == [["prem@example.com"]]

    assert anon.get(f"/api/sessions/{token}/").status_code == 200
    session.refresh_from_db()
    assert session.events.filter(event_type="signing.viewed").count() == 1
    assert len([m for m in mail.outbox if m.subject.startswith("Viewed:")]) == 1


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_viewed_alert_suppressed_when_opted_out_free_or_ownerless():
    from django.core import mail

    # Premium but opted out: the view is still recorded, no email.
    opted = _make_account(email="optout@example.com", tier=Account.Tier.PREMIUM)
    opted.alerts_enabled = False
    opted.save(update_fields=["alerts_enabled"])
    _, token1 = _owned_session(_auth_client(opted), opted.email)

    # Free-tier owner: no email.
    free = _make_account(email="free@example.com", tier=Account.Tier.FREE)
    _, token2 = _owned_session(_auth_client(free), free.email)

    # Ownerless document: no email.
    anon = APIClient()
    _, token3 = _owned_session(anon, "nobody@example.com")

    for token in (token1, token2, token3):
        assert anon.get(f"/api/sessions/{token}/").status_code == 200
        session = SigningSession.objects.get(token=token)
        assert session.viewed_at is not None
        assert session.events.filter(event_type="signing.viewed").count() == 1
    # No VIEWED alerts fired (the invites in the outbox are expected).
    assert [m for m in mail.outbox if m.subject.startswith("Viewed:")] == []


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_expiry_alert_command_sends_stamps_and_does_not_resend():
    from django.core import mail
    from django.core.management import call_command

    account = _make_account(email="prem2@example.com", tier=Account.Tier.PREMIUM)
    _, token = _owned_session(_auth_client(account), account.email)
    session = SigningSession.objects.get(token=token)
    session.expires_at = timezone.now() + timedelta(hours=12)
    session.save(update_fields=["expires_at"])

    # A second session expiring well outside the window stays quiet.
    _, far_token = _owned_session(_auth_client(account), account.email)
    far = SigningSession.objects.get(token=far_token)
    far.expires_at = timezone.now() + timedelta(hours=48)
    far.save(update_fields=["expires_at"])

    call_command("send_expiry_alerts")
    session.refresh_from_db()
    far.refresh_from_db()
    assert session.expiry_alert_sent_at is not None
    assert far.expiry_alert_sent_at is None
    alerts = [m for m in mail.outbox if m.subject.startswith("Expiring soon:")]
    assert len(alerts) == 1
    assert alerts[0].to == ["prem2@example.com"]
    assert "/d/" in alerts[0].body

    call_command("send_expiry_alerts")
    assert len([m for m in mail.outbox if m.subject.startswith("Expiring soon:")]) == 1
# Issue #15: signing link + QR for pending documents. The dashboard and
# the tracking page need the active session's signing URL — but the
# token is a credential, so it is exposed to the document owner only.
# ---------------------------------------------------------------------------


def _owned_doc_with_session(account):
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    session = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Jane", "signer_email": "jane@example.com"}, format="json"
    ).json()
    return client, doc_id, session


@pytest.mark.django_db
def test_dashboard_includes_signing_url_for_pending_owned_document():
    account = _make_account(email="dash@example.com")
    client, doc_id, session = _owned_doc_with_session(account)
    body = client.get("/api/documents/mine/").json()
    (doc,) = body["documents"]
    assert doc["id"] == doc_id
    assert doc["signing_url"] == session["signing_url"]
    assert doc["signing_url"].endswith(f"/s/{session['token']}")


@pytest.mark.django_db
def test_dashboard_signing_url_is_null_without_usable_session():
    # No session yet.
    account = _make_account(email="nosession@example.com")
    client = _auth_client(account)
    upload_document(client, email=account.email)
    (doc,) = client.get("/api/documents/mine/").json()["documents"]
    assert doc["signing_url"] is None

    # Session exists but expired.
    account2 = _make_account(email="expired@example.com")
    client2, _, session2 = _owned_doc_with_session(account2)
    SigningSession.objects.filter(pk=session2["id"]).update(
        expires_at=timezone.now() - timedelta(days=1)
    )
    (doc2,) = client2.get("/api/documents/mine/").json()["documents"]
    assert doc2["signing_url"] is None

    # Session already signed.
    account3 = _make_account(email="signed3@example.com")
    client3, _, session3 = _owned_doc_with_session(account3)
    assert APIClient().post(
        f"/api/sessions/{session3['token']}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    ).status_code == 200
    (doc3,) = client3.get("/api/documents/mine/").json()["documents"]
    assert doc3["signing_url"] is None


@pytest.mark.django_db
def test_document_view_exposes_signing_url_only_to_owner():
    account = _make_account(email="viewer@example.com")
    client, doc_id, session = _owned_doc_with_session(account)
    token = session["token"]

    # Owner: signing_url present and pointing at the session token.
    owner_body = client.get(f"/api/documents/{doc_id}/").json()
    assert owner_body["signing_url"] == session["signing_url"]
    assert owner_body["signing_url"].endswith(f"/s/{token}")

    # Anonymous: no signing URL field, and the token leaks nowhere.
    anon_response = APIClient().get(f"/api/documents/{doc_id}/")
    assert anon_response.status_code == 200
    assert "signing_url" not in anon_response.json()
    assert token not in anon_response.content.decode()

    # Stale-token requester degrades to anonymous (OptionalToken auth).
    stale_response = _stale_client().get(f"/api/documents/{doc_id}/")
    assert stale_response.status_code == 200
    assert "signing_url" not in stale_response.json()

    # A different authenticated account is not the owner.
    other = _auth_client(_make_account(email="other@example.com"))
    other_body = other.get(f"/api/documents/{doc_id}/").json()
    assert "signing_url" not in other_body

# ---------------------------------------------------------------------------
# Multi-signer flow (premium, strictly sequential): the document seals
# once, only when every signer in order has signed.
# ---------------------------------------------------------------------------


def _premium_multi_doc(names=("Signer One", "Signer Two", "Signer Three")):
    account = _make_account(
        email="multi@example.com", tier=Account.Tier.PREMIUM
    )
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    tokens = []
    for name in names:
        resp = client.post(
            f"/api/documents/{doc_id}/sessions/",
            {
                "signer_name": name,
                "signer_email": f"{name.lower().replace(' ', '.')}@example.com",
            },
            format="json",
        )
        assert resp.status_code == 201, resp.content
        tokens.append(resp.json()["token"])
    return client, doc_id, tokens


def _sign(token, client=None):
    return (client or APIClient()).post(
        f"/api/sessions/{token}/sign/",
        {"signature": SIG_DATA_URL, "device_info": {}},
        format="json",
    )


@pytest.mark.django_db
def test_premium_sessions_get_sequential_orders():
    _, doc_id, _ = _premium_multi_doc()
    sessions = list(
        SigningSession.objects.filter(document_id=doc_id).order_by("order")
    )
    assert [s.order for s in sessions] == [0, 1, 2]
    assert [s.signer_name for s in sessions] == [
        "Signer One",
        "Signer Two",
        "Signer Three",
    ]


@pytest.mark.django_db
def test_free_account_document_still_rejects_second_session():
    account = _make_account(email="free-acct@example.com")
    client = _auth_client(account)
    doc_id = upload_document(client, email=account.email).json()["id"]
    assert client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "One", "signer_email": "one@example.com"}, format="json"
    ).status_code == 201
    second = client.post(
        f"/api/documents/{doc_id}/sessions/", {"signer_name": "Two", "signer_email": "two@example.com"}, format="json"
    )
    assert second.status_code == 409


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_three_signer_document_seals_once_in_order():
    from django.core import mail

    _, doc_id, tokens = _premium_multi_doc()
    anon = APIClient()

    # Signer 2's link is not active until signer 1 signs.
    detail = anon.get(f"/api/sessions/{tokens[1]}/").json()
    assert detail["usable"] is False
    assert detail["waiting"] is True
    assert _sign(tokens[1], anon).status_code == 410

    # Signer 1 signs: recorded, but the document is NOT final.
    first = _sign(tokens[0], anon)
    assert first.status_code == 200
    assert first.json()["status"] == "signed"
    assert first.json()["document_final"] is False
    doc = Document.objects.get(pk=doc_id)
    assert doc.status == Document.Status.PENDING
    assert not doc.signed_file
    assert anon.get(f"/api/documents/{doc_id}/").json()["signed_file"] is None
    # The per-session document endpoint still serves the ORIGINAL pdf —
    # nothing partially signed is downloadable as final.
    served = anon.get(f"/api/sessions/{tokens[0]}/document/")
    assert served.status_code == 200
    assert b"SIGNATURE CERTIFICATE" not in b"".join(served.streaming_content)
    assert [m for m in mail.outbox if m.subject.startswith("Signed:")] == []  # no seal notifications yet (invites may exist)

    # Signer 2 activates and signs; still not final.
    assert anon.get(f"/api/sessions/{tokens[1]}/").json()["usable"] is True
    second = _sign(tokens[1], anon)
    assert second.json()["document_final"] is False
    doc.refresh_from_db()
    assert doc.status == Document.Status.PENDING
    assert not doc.signed_file
    # No seal notification yet (a viewed alert may exist: signer 2's
    # session was viewed once it activated — alerts are issue #9).
    assert [m for m in mail.outbox if m.subject.startswith("Signed:")] == []

    # Signer 3 seals the document, exactly once.
    third = _sign(tokens[2], anon)
    assert third.status_code == 200
    assert third.json()["document_final"] is True
    doc.refresh_from_db()
    assert doc.status == Document.Status.SIGNED
    assert doc.signed_file
    seal_mail = [m for m in mail.outbox if m.subject.startswith("Signed:")]
    # Sender copy plus one copy per signer (every session now carries an
    # email under email-first delivery), all at final seal only.
    assert len(seal_mail) == 4
    assert seal_mail[0].to == ["multi@example.com"]
    assert sorted(m.to[0] for m in seal_mail[1:]) == [
        "signer.one@example.com",
        "signer.three@example.com",
        "signer.two@example.com",
    ]

    # The certificate lists all three signers.
    pdf = pymupdf.open(stream=doc.signed_file.read())
    text = "".join(page.get_text() for page in pdf)
    for name in ("Signer One", "Signer Two", "Signer Three"):
        assert name in text

    # Audit: three signature.submitted events in signing order, one
    # document.sealed, and every session chain verifies.
    submitted = [
        e
        for e in AuditEvent.objects.filter(
            session__document_id=doc_id
        ).order_by("created_at")
        if e.event_type == "signature.submitted"
    ]
    assert [e.session.signer_name for e in submitted] == [
        "Signer One",
        "Signer Two",
        "Signer Three",
    ]
    sealed = AuditEvent.objects.filter(
        session__document_id=doc_id, event_type="document.sealed"
    )
    assert sealed.count() == 1
    for session in SigningSession.objects.filter(document_id=doc_id):
        assert verify_chain(session) is True

    # Document payload carries the per-signer state for the sender.
    detail = anon.get(f"/api/documents/{doc_id}/").json()
    assert [(s["signer_name"], s["status"], s["order"]) for s in detail["sessions"]] == [
        ("Signer One", "signed", 0),
        ("Signer Two", "signed", 1),
        ("Signer Three", "signed", 2),
    ]


# ---------------------------------------------------------------------------
# Complex documents (#13): sender-placed fields, per-signer assignment,
# sequential vs parallel signing, field stamping at seal, email invites.
# ---------------------------------------------------------------------------

from .models import DocumentField


def make_multipage_pdf_bytes(pages=2) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {i + 1} body")
    return doc.tobytes()


def upload_mode_document(client, email, mode=None, pages=2):
    pdf = SimpleUploadedFile(
        "complex.pdf", make_multipage_pdf_bytes(pages), content_type="application/pdf"
    )
    data = {"title": "Complex", "sender_email": email, "file": pdf}
    if mode is not None:
        data["signing_mode"] = mode
    return client.post("/api/documents/", data, format="multipart")


def _make_session(client, doc_id, name, email):
    resp = client.post(
        f"/api/documents/{doc_id}/sessions/",
        {"signer_name": name, "signer_email": email},
        format="json",
    )
    assert resp.status_code == 201, resp.content
    return resp.json()


def _place(client, doc_id, session_id, field_type, page=1, x=0.1, y=0.2, required=True):
    resp = client.post(
        f"/api/documents/{doc_id}/fields/",
        {
            "session_id": session_id,
            "type": field_type,
            "page": page,
            "x": x,
            "y": y,
            "required": required,
        },
        format="json",
    )
    assert resp.status_code == 201, resp.content
    return resp.json()


def _sign_with_fields(token, field_values):
    return APIClient().post(
        f"/api/sessions/{token}/sign/",
        {
            "signature": SIG_DATA_URL,
            "device_info": {},
            "field_values": field_values,
        },
        format="json",
    )


def _premium_fields_setup(mode="sequential"):
    account = _make_account(
        email=f"fields-{mode}@example.com", tier=Account.Tier.PREMIUM
    )
    client = _auth_client(account)
    resp = upload_mode_document(client, account.email, mode=mode)
    assert resp.status_code == 201, resp.content
    doc_id = resp.json()["id"]
    alice = _make_session(client, doc_id, "Alice", "alice@example.com")
    bob = _make_session(client, doc_id, "Bob", "bob@example.com")
    return client, doc_id, alice, bob


@pytest.mark.django_db
def test_signing_mode_defaults_sequential_and_validates():
    client = APIClient()
    ok = upload_mode_document(client, "mode-default@example.com")
    assert ok.status_code == 201
    assert ok.json()["signing_mode"] == "sequential"
    bad = upload_mode_document(client, "mode-bad@example.com", mode="sideways")
    assert bad.status_code == 400


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_complex_document_fields_end_to_end_sequential():
    from django.core import mail

    client, doc_id, alice, bob = _premium_fields_setup("sequential")

    # Invite delivery: Alice (active) was invited at creation; Bob is
    # held by the sequence and has not been invited yet.
    invited = [m.to[0] for m in mail.outbox if m.subject.startswith("Please sign")]
    assert invited == ["alice@example.com"]

    # Fields: Alice signs/text/dates, Bob initials + checkbox.
    a_sig = _place(client, doc_id, alice["id"], "signature", page=1, x=0.1, y=0.2)
    a_text = _place(client, doc_id, alice["id"], "text", page=1, x=0.1, y=0.45)
    a_date = _place(client, doc_id, alice["id"], "date", page=2, x=0.1, y=0.2)
    b_init = _place(client, doc_id, bob["id"], "initials", page=1, x=0.55, y=0.2)
    b_check = _place(client, doc_id, bob["id"], "checkbox", page=2, x=0.1, y=0.5)

    # The ceremony payload carries only the signer's own fields.
    payload = APIClient().get(f"/api/sessions/{alice['token']}/").json()
    assert [f["id"] for f in payload["fields"]] == [a_sig["id"], a_text["id"], a_date["id"]]
    assert [f["type"] for f in payload["fields"]] == ["signature", "text", "date"]

    # Signer 2 cannot submit before signer 1 (sequential gating).
    assert _sign(bob["token"]).status_code == 410

    # Missing required fields -> 400, and completing another signer's
    # field -> 400.
    assert _sign_with_fields(alice["token"], []).status_code == 400
    partial = [
        {"field_id": a_sig["id"], "value": SIG_DATA_URL},
        {"field_id": a_text["id"], "value": "Alice agreed"},
    ]
    assert _sign_with_fields(alice["token"], partial).status_code == 400
    foreign = partial + [
        {"field_id": a_date["id"], "value": "2026-10-20"},
        {"field_id": b_check["id"], "value": True},
    ]
    assert _sign_with_fields(alice["token"], foreign).status_code == 400

    # Alice completes all of her fields.
    done = _sign_with_fields(
        alice["token"],
        partial + [{"field_id": a_date["id"], "value": "2026-10-20"}],
    )
    assert done.status_code == 200, done.content
    assert done.json()["document_final"] is False
    doc = Document.objects.get(pk=doc_id)
    assert doc.status == Document.Status.PENDING

    # Bob was invited when Alice's signature activated his session.
    bob_invites = [
        m for m in mail.outbox
        if m.to == ["bob@example.com"] and m.subject.startswith("Please sign")
    ]
    assert len(bob_invites) == 1
    assert f"/s/{bob['token']}" in bob_invites[0].body

    # Bob completes his fields; the document seals exactly once.
    final = _sign_with_fields(
        bob["token"],
        [
            {"field_id": b_init["id"], "value": SIG_DATA_URL},
            {"field_id": b_check["id"], "value": True},
        ],
    )
    assert final.status_code == 200, final.content
    assert final.json()["document_final"] is True
    doc.refresh_from_db()
    assert doc.status == Document.Status.SIGNED
    assert DocumentField.objects.filter(
        document_id=doc_id, completed_at__isnull=True
    ).count() == 0

    # Sealed PDF: field values are stamped on their pages.
    pdf = pymupdf.open(stream=doc.signed_file.read())
    page1 = pdf[0].get_text()
    page2 = pdf[1].get_text()
    assert "Alice agreed" in page1
    assert "2026-10-20" in page2
    assert "X" in page2  # the checkbox mark

    # field.completed events in both chains; every chain verifies.
    for session_id, count in ((alice["id"], 3), (bob["id"], 2)):
        session = SigningSession.objects.get(pk=session_id)
        events = session.events.filter(event_type="field.completed")
        assert events.count() == count
        assert verify_chain(session) is True


@pytest.mark.django_db
@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
def test_parallel_mode_lets_signer_two_go_first():
    from django.core import mail

    client, doc_id, alice, bob = _premium_fields_setup("parallel")

    # Parallel: both signers were invited at creation.
    invited = {m.to[0] for m in mail.outbox if m.subject.startswith("Please sign")}
    assert invited == {"alice@example.com", "bob@example.com"}

    a_text = _place(client, doc_id, alice["id"], "text")
    b_text = _place(client, doc_id, bob["id"], "text")
    anon = APIClient()
    assert anon.get(f"/api/sessions/{bob['token']}/").json()["usable"] is True

    # Bob (order 1) signs first — allowed in parallel mode.
    first = _sign_with_fields(
        bob["token"], [{"field_id": b_text["id"], "value": "Bob first"}]
    )
    assert first.status_code == 200, first.content
    assert first.json()["document_final"] is False

    second = _sign_with_fields(
        alice["token"], [{"field_id": a_text["id"], "value": "Alice second"}]
    )
    assert second.json()["document_final"] is True
    doc = Document.objects.get(pk=doc_id)
    pdf = pymupdf.open(stream=doc.signed_file.read())
    text = "".join(page.get_text() for page in pdf)
    assert "Bob first" in text and "Alice second" in text


@pytest.mark.django_db
def test_field_placement_rejects_non_owner_and_bad_values():
    client, doc_id, alice, _ = _premium_fields_setup("sequential")

    # Unauthenticated and non-owner access is refused.
    assert APIClient().get(f"/api/documents/{doc_id}/fields/").status_code == 401
    other = _auth_client(_make_account(email="intruder@example.com"))
    assert other.get(f"/api/documents/{doc_id}/fields/").status_code == 403
    assert other.post(
        f"/api/documents/{doc_id}/fields/",
        {"session_id": alice["id"], "type": "text", "page": 1, "x": 0.1, "y": 0.1},
        format="json",
    ).status_code == 403
    assert other.get(f"/api/documents/{doc_id}/page/1.png").status_code == 403

    def place(**overrides):
        body = {
            "session_id": alice["id"],
            "type": "text",
            "page": 1,
            "x": 0.1,
            "y": 0.1,
        }
        body.update(overrides)
        return client.post(f"/api/documents/{doc_id}/fields/", body, format="json")

    assert place(page=99).status_code == 400  # beyond the 2-page document
    assert place(page=0).status_code == 400
    assert place(x=1.5).status_code == 400
    assert place(y=-0.1).status_code == 400
    assert place(type="bogus").status_code == 400
    assert place(session_id="00000000-0000-0000-0000-000000000000").status_code == 400

    # Page rendering: owner gets a PNG; out-of-range is a 404.
    png = client.get(f"/api/documents/{doc_id}/page/1.png")
    assert png.status_code == 200
    assert png["Content-Type"] == "image/png"
    assert png.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert client.get(f"/api/documents/{doc_id}/page/99.png").status_code == 404

    # Signer page render: the active signer can view; nobody else
    # holds a token for the waiting signer here, but the waiting
    # session's own render is refused.
    anon = APIClient()
    assert anon.get(f"/api/sessions/{alice['token']}/page/1.png").status_code == 200
    sessions = list(SigningSession.objects.filter(document_id=doc_id).order_by("order"))
    assert anon.get(f"/api/sessions/{sessions[1].token}/page/1.png").status_code == 410

    # Delete works for the owner while the field is uncompleted.
    field = _place(client, doc_id, alice["id"], "text")
    assert client.delete(f"/api/documents/{doc_id}/fields/{field['id']}/").status_code == 204
    listing = client.get(f"/api/documents/{doc_id}/fields/").json()
    assert listing["page_count"] == 2
    assert listing["fields"] == []
